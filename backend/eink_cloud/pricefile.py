"""Prijslijsten inlezen: CSV of Excel uit een kassa, weegschaalsoftware of door de winkel zelf gemaakt.

Werkwijze: bestand lezen → kolommen herkennen (of door de gebruiker laten kiezen) → per regel een
product maken → eerst een voorbeeld tonen met fouten per regel → pas daarna importeren.
De gekozen kolomindeling wordt per winkel bewaard, zodat een vaste export (bv. elke ochtend uit de
weegschaal) zonder tussenkomst ingelezen kan worden.

Merkspecifieke formaten (bv. XML van Mettler Toledo RetailSuite) komen als eigen lezer hierbij,
zodra er voorbeeldbestanden zijn; ze leveren dezelfde tabel (kolommen + regels) op.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from pydantic import ValidationError

from .schemas import ProductBatchItem

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000

# Productveld → herkenbare kolomnamen (kleine letters, zonder leestekens).
FIELDS = {
    "sku": ("plu", "plunr", "plunummer", "artikelnummer", "artnr", "artikelnr", "artikelcode", "sku", "code",
            "itemnumber", "itemno", "barcode", "ean", "nummer", "nr"),
    "name": ("naam", "omschrijving", "artikelomschrijving", "artikelnaam", "artikel", "product", "productnaam",
             "description", "name", "tekst", "text", "benaming"),
    "price": ("prijs", "verkoopprijs", "vkprijs", "vkp", "price", "prijsinclbtw", "prijsincl", "eenheidsprijs",
              "prijsperkg", "kiloprijs", "consumentenprijs", "bedrag"),
    "unit": ("eenheid", "unit", "verkoopeenheid", "per", "prijseenheid", "weegartikel", "gewogen"),
    "unit_price": ("prijsperkgvoorverpakt", "kgprijs", "literprijs", "prijsperliter", "unitprice", "eenheidprijs"),
    "origin": ("herkomst", "land", "oorsprong", "origin", "landvanherkomst", "afkomst"),
    "promo_text": ("actie", "actietekst", "promotie", "promo", "aanbieding", "promotext"),
    "was_price": ("vanprijs", "van", "oudeprijs", "adviesprijs", "normaleprijs", "wasprice", "regularprice"),
    "description": ("beschrijving", "toelichting", "ingredienten", "ingrediënten", "allergenen", "info", "details",
                    "latijnsenaam", "wetenschappelijkenaam"),
    "template": ("ontwerp", "sjabloon", "template", "layout"),
    "options": ("sauzen", "saus", "keuze", "keuzes", "opties", "options"),
}
FIELD_LABELS = {
    "sku": "Artikelnummer / PLU", "name": "Naam", "price": "Prijs", "unit": "Eenheid",
    "unit_price": "Prijs per kg (voorverpakt)", "origin": "Herkomst", "promo_text": "Actietekst",
    "was_price": "Van-prijs", "description": "Omschrijving", "template": "Ontwerp", "options": "Keuzes (sauzen)",
}
REQUIRED = ("sku", "name", "price")
UNIT_ALIASES = {
    "kg": "kg", "kilo": "kg", "perkg": "kg", "kilogram": "kg", "gewicht": "kg", "w": "kg", "weeg": "kg",
    "gewogen": "kg", "ja": "kg", "j": "kg", "1": "kg", "true": "kg",
    "st": "st", "stuk": "st", "stuks": "st", "perstuk": "st", "stk": "st", "pc": "st", "pcs": "st", "s": "st",
    "nee": "st", "n": "st", "0": "st", "false": "st", "": "st",
    "100g": "100g", "100gr": "100g", "per100g": "100g", "ons": "100g",
    "l": "l", "liter": "l", "ltr": "l", "pak": "pak", "verpakking": "pak",
}


class PriceFileError(ValueError):
    pass


@dataclass
class Table:
    columns: list[str]
    rows: list[list[str]]


@dataclass
class ImportResult:
    columns: list[str]
    mapping: dict[str, str]
    rows_total: int
    items: list[ProductBatchItem] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)  # {"row": regelnummer in het bestand, "message": ...}
    skipped: int = 0


def _key(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):  # Excel op Windows bewaart CSV vaak als cp1252
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise PriceFileError("tekenset van het bestand niet herkend")


def read_table(data: bytes, filename: str) -> Table:
    if len(data) > MAX_BYTES:
        raise PriceFileError(f"bestand is te groot (max {MAX_BYTES // 1024 // 1024} MB)")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        rows = _read_xlsx(data)
    elif name.endswith(".xls"):
        raise PriceFileError("oud Excel-formaat (.xls): sla het bestand op als .xlsx of .csv")
    else:
        rows = _read_csv(_decode(data))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise PriceFileError("het bestand is leeg")
    if len(rows) > MAX_ROWS + 1:
        raise PriceFileError(f"te veel regels (max {MAX_ROWS})")
    header = [c.strip() or f"kolom {i + 1}" for i, c in enumerate(rows[0])]
    width = len(header)
    body = [(r + [""] * width)[:width] for r in rows[1:]]
    return Table(header, body)


def _read_csv(text: str) -> list[list[str]]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - openpyxl staat in requirements
        raise PriceFileError("Excel-bestanden worden niet ondersteund op deze server") from None
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # beschadigd of geen Excel
        raise PriceFileError(f"Excel-bestand kan niet gelezen worden: {exc}") from None
    sheet = wb.worksheets[0]
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([_cell(v) for v in row])
        if len(rows) > MAX_ROWS + 1:
            break
    wb.close()
    return rows


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        # 12.95 blijft "12.95"; 1001.0 (artikelnummer) wordt "1001"
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def guess_mapping(columns: list[str]) -> dict[str, str]:
    """Herkent kolommen op naam; elke kolom hoogstens één keer."""
    mapping: dict[str, str] = {}
    keys = {c: _key(c) for c in columns}
    for fieldname, names in FIELDS.items():
        for name in names:  # volgorde = voorkeur
            match = next((c for c, k in keys.items() if k == name and c not in mapping.values()), None)
            if match:
                mapping[fieldname] = match
                break
    return mapping


def parse_price(text: str) -> int:
    t = text.replace("€", "").replace("EUR", "").replace(" ", "").replace(" ", "").strip()
    if not t:
        raise ValueError("prijs ontbreekt")
    if "," in t and "." in t:  # 1.234,56 of 1,234.56
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        value = Decimal(t)
    except InvalidOperation:
        raise ValueError(f"ongeldige prijs '{text}'") from None
    if value < 0:
        raise ValueError(f"negatieve prijs '{text}'")
    return int((value * 100).quantize(Decimal("1")))


def parse_unit(text: str) -> str:
    unit = UNIT_ALIASES.get(_key(text))
    if unit is None:
        raise ValueError(f"onbekende eenheid '{text}' (gebruik st, kg, 100g, l of pak)")
    return unit


def build_items(table: Table, mapping: dict[str, str], default_unit: str = "st") -> ImportResult:
    unknown = [c for c in mapping.values() if c and c not in table.columns]
    if unknown:
        raise PriceFileError(f"kolom(men) niet in het bestand: {', '.join(unknown)}")
    missing = [FIELD_LABELS[f] for f in REQUIRED if not mapping.get(f)]
    result = ImportResult(columns=table.columns, mapping={k: v for k, v in mapping.items() if v}, rows_total=len(table.rows))
    if missing:
        result.errors.append({"row": None, "message": f"kies een kolom voor: {', '.join(missing)}"})
        return result
    idx = {f: table.columns.index(c) for f, c in result.mapping.items()}
    seen: dict[str, int] = {}
    for n, row in enumerate(table.rows, start=2):  # regel 1 = kolomnamen
        def get(f: str) -> str:
            return row[idx[f]].strip() if f in idx else ""

        sku = get("sku")
        if not sku and not get("name") and not get("price"):
            result.skipped += 1
            continue
        try:
            if not sku:
                raise ValueError("artikelnummer ontbreekt")
            if sku in seen:
                raise ValueError(f"artikelnummer {sku} komt dubbel voor (ook op regel {seen[sku]})")
            unit = parse_unit(get("unit")) if "unit" in idx and get("unit") else default_unit
            unit_price = parse_price(get("unit_price")) if get("unit_price") else None
            was_price = parse_price(get("was_price")) if get("was_price") else None
            item = ProductBatchItem(sku=sku, name=get("name"), price_cents=parse_price(get("price")), unit=unit,
                                    unit_price_cents=unit_price, origin=get("origin") or None,
                                    promo_text=get("promo_text") or None, was_price_cents=was_price,
                                    description=get("description") or None, template=get("template").lower() or None,
                                    options=get("options") or None)
            if not item.name:
                raise ValueError("naam ontbreekt")
        except ValidationError as exc:
            result.errors.append({"row": n, "message": "; ".join(e["msg"] for e in exc.errors())})
            continue
        except ValueError as exc:
            result.errors.append({"row": n, "message": str(exc)})
            continue
        seen[sku] = n
        result.items.append(item)
    return result
